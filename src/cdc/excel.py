"""Export the master sheet to .xlsx (one tab per subsystem) and import edited inputs back."""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from cdc.engine import CATALOGUE
from cdc.models import SUBSYSTEMS
from cdc.store import MasterSheet

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="E8F0FE")
STATUS_FILL = {"PASS": "C6EFCE", "MARGINAL": "FFEB9C", "FAIL": "FFC7CE", "NO DATA": "D9D9D9"}
COLS = ["Parameter", "Kind", "Value", "Unit", "Description", "Group / Source"]


def _sheet(ws, header, rows):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in rows:
        ws.append(r)
    for i, _ in enumerate(header, 1):
        width = max(len(str(ws.cell(row=j, column=i).value or "")) for j in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = "A2"


def export_design(ms: MasterSheet, design: str) -> bytes:
    vals = ms.results(design).values
    wb = Workbook()
    master = wb.active
    master.title = "Master"
    rows = []
    for name, meta in CATALOGUE.items():
        p = meta["param"]
        rows.append([name, meta["kind"], vals.get(name), p.unit, p.desc, p.group or p.source])
    _sheet(master, COLS, rows)
    for r in range(2, master.max_row + 1):
        if master.cell(row=r, column=2).value == "input":
            master.cell(row=r, column=3).fill = INPUT_FILL

    for s in SUBSYSTEMS:
        ws = wb.create_sheet(s.title.replace("/", "-").translate({ord(c): None for c in "\\?*[]:"})[:31])
        rows = [[s.full(p.key), "input", vals[s.full(p.key)], p.unit, p.desc, p.source] for p in s.inputs]
        rows += [[f"{s.name}.{k}", f"get ← {src}", vals[src], CATALOGUE[src]["param"].unit,
                  CATALOGUE[src]["param"].desc, ""] for k, src in s.gets.items()]
        rows += [[s.full(p.key), "send →", vals[s.full(p.key)], p.unit, p.desc, ""] for p in s.sends]
        _sheet(ws, COLS, rows)

    comp = wb.create_sheet("Comparison")
    rows = [[c["requirement"], c["parameter"], c["accepted"], c["required"], c["value"], c["unit"],
             c["status"], None if c["margin_pct"] is None else round(c["margin_pct"], 2), c["source"]]
            for c in ms.compliance(design)]
    _sheet(comp, ["Requirement", "Parameter", "Accepted if", "Required", "Value", "Unit", "Status", "Margin %",
                  "Source"], rows)
    for r in range(2, comp.max_row + 1):
        st = comp.cell(row=r, column=7)
        st.fill = PatternFill("solid", fgColor=STATUS_FILL.get(st.value, "FFFFFF"))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_inputs(data: bytes) -> dict[str, float]:
    """Read edited input values from the 'Master' tab of an exported workbook."""
    ws = load_workbook(io.BytesIO(data), data_only=True)["Master"]
    out = {}
    for name, kind, value, *_ in ws.iter_rows(min_row=2, values_only=True):
        if kind == "input" and name in CATALOGUE and isinstance(value, (int, float)):
            out[name] = float(value)
    return out
